"""Tests for resume file parsing."""

from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from ats_checker.parser import extract_text, split_resume_sections


def test_extract_docx(sample_docx: Path):
    text = extract_text(sample_docx)
    assert "Jane Doe" in text
    assert "jane@example.com" in text


def test_unsupported_format(tmp_path: Path):
    bad = tmp_path / "resume.txt"
    bad.write_text("hello")
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(bad)


def test_docx_suffix_case(sample_docx: Path):
    upper = sample_docx.parent / "RESUME.DOCX"
    sample_docx.rename(upper)
    text = extract_text(upper)
    assert "Jane Doe" in text


def test_split_resume_sections_detects_common_headings(good_resume_text: str):
    sections = split_resume_sections(good_resume_text)
    names = [name for name, _ in sections]

    assert names[:2] == ["Header", "Professional Summary"]
    assert "Experience" in names
    assert "Education" in names
    assert "Skills" in names

    section_map = dict(sections)
    assert "John Doe" in section_map["Header"]
    assert "Led a team of 8 engineers" in section_map["Experience"]


def test_split_resume_sections_falls_back_to_single_block(bare_resume_text: str):
    sections = split_resume_sections(bare_resume_text)
    assert sections == [("Resume", bare_resume_text)]


def test_docx_preserves_headers_tables_order_and_native_bullets(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "jane@example.com"
    doc.sections[0].footer.paragraphs[0].text = "Portfolio: example.com"
    doc.add_paragraph("Experience")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Reduced latency by 30%."
    table.cell(1, 0).text = "Skills"
    table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Python, SQL"
    style = doc.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["List Bullet"]
    doc.add_paragraph("Delivered APIs", style=style)
    plain = doc.add_paragraph("Education", style=style)
    plain._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    doc.add_section()
    doc.add_paragraph("BS Computer Science")
    path = tmp_path / "resume.docx"
    doc.save(path)

    assert extract_text(path).splitlines() == [
        "jane@example.com", "Experience", "Reduced latency by 30%.",
        "Skills", "Python, SQL", "- Delivered APIs", "Education",
        "BS Computer Science", "Portfolio: example.com",
    ]


def test_table_only_docx_retains_resume_text(tmp_path):
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Jane Doe jane@example.com"
    path = tmp_path / "table-only.docx"
    doc.save(path)
    assert extract_text(path) == "Jane Doe jane@example.com"
